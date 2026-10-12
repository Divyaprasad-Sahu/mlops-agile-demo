"""Build the final lab report using bundled python-docx, with actual evidence."""
import json
import ast
from pathlib import Path
import subprocess

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
metrics = json.loads((ROOT / "models/metrics.json").read_text())
evidence_file = ROOT / "docs/evidence/github_evidence.json"
evidence = json.loads(evidence_file.read_text()) if evidence_file.exists() else {}
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.7)
for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(7)
doc.styles["Normal"].font.size = Pt(11)
doc.styles["Normal"].paragraph_format.line_spacing = 1.08
doc.styles["Title"].font.size = Pt(26)
doc.styles["Heading 1"].font.size = Pt(18)
doc.styles["Heading 2"].font.size = Pt(13)
for style in doc.styles:
    for border in style.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
doc.core_properties.title = "Agile and Scrum MLOps Experiment"
doc.core_properties.author = "Divyaprasad Sahu"


def p(text, style=None):
    return doc.add_paragraph(text, style)


def h(text):
    doc.add_heading(text, 2)


def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = str(text)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri, row in enumerate(t.rows):
        keep = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(keep)
        for ci, cell in enumerate(row.cells):
            if widths:
                cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + edge)
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn("w:" + key), value)
                borders.append(border)
            props.append(borders)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                item = OxmlElement("w:" + side)
                item.set(qn("w:w"), "90")
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            props.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "17365D" if ri == 0 else ("F1F4F8" if ri % 2 == 0 else "FFFFFF"))
            props.append(shading)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    run.font.size = Pt(9.5)
                    if ri == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    p("")
    return t


def code(text):
    para = p(text)
    para.paragraph_format.line_spacing = 1
    for run in para.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(9)


def source_excerpt(relative_path, start, end):
    """Copy actual source lines with editor-style line numbers into Word."""
    lines = (ROOT / relative_path).read_text(encoding="utf-8").splitlines()
    p(f"{relative_path}   |   Lines {start} to {end}")
    for number in range(start, end + 1):
        para = p("")
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.05
        number_run = para.add_run(f"{number:>3}  ")
        number_run.font.name = "Consolas"
        number_run.font.size = Pt(9)
        number_run.font.color.rgb = RGBColor.from_string("666666")
        run = para.add_run(lines[number - 1])
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        if lines[number - 1].lstrip().startswith("#"):
            run.font.color.rgb = RGBColor.from_string("38761D")
        elif lines[number - 1].lstrip().startswith(("def ", "if ", "return ", "with ")):
            run.font.color.rgb = RGBColor.from_string("17365D")
    p("")


def function_excerpt(relative_path, function_name):
    source = (ROOT / relative_path).read_text(encoding="utf-8")
    node = next(node for node in ast.parse(source).body
                if isinstance(node, ast.FunctionDef) and node.name == function_name)
    source_excerpt(relative_path, node.lineno, node.end_lineno)


def link(label, url):
    para = p("")
    rel = para.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = label + "  " + url
    run.append(text)
    element.append(run)
    para._p.append(element)


p("Experiment 1", "Subtitle")
p("Agile and Scrum MLOps Experiment", "Title")
p("Iris flower classification with Git and automated delivery", "Subtitle")
p("Divyaprasad Sahu")
p("Roll number  ____________________")
p("Institution  ____________________")
p("Course  ________________________")
p("Faculty  _______________________")
p("Experiment date  12 October 2026")
h("Aim")
p("Demonstrate Agile methodology and Scrum through an Iris machine learning project, including roles, backlogs, Sprint ceremonies, a working model, version control, automated testing and continuous delivery.")
h("Learning objectives")
for text in ("Explain Agile principles and Scrum responsibilities.", "Prioritise a Product Backlog and plan a one-week Sprint.", "Simulate Daily Scrum, Sprint Review and Sprint Retrospective.", "Train, evaluate and save a reproducible Decision Tree classifier.", "Use Git and GitHub to manage changes and demonstrate automated test and delivery workflows."):
    p(text, "List Bullet")
h("Outcome")
p(f"The classifier achieved {metrics['accuracy']:.0%} accuracy on the fixed 30-row held-out split. The final acceptance suite contains 19 passing tests. Scrum ceremonies are an individual educational simulation; model results and GitHub workflow evidence are actual executions.")

page("Principles and project preparation")
h("Agile and Scrum theory")
p("Agile develops useful software in small increments and adapts through feedback. Scrum organises this work around a Product Goal, a Product Backlog, timeboxed Sprints, a Sprint Backlog and a usable Increment. Inspection and adaptation occur in planning, daily coordination, review and retrospective.")
p("Product Owner, Scrum Master and Developers are Scrum accountabilities. The handout's ML Engineer, Data Engineer and Tester roles are specialisations within Developers. Daily Scrum is a Developers event; this report retains all-role updates to match the course simulation.")
h("Step 1 Create the workspace")
p("The standalone MLOps_Agile_Demo folder contains source, data, models, tests and documentation. The handout's example drive location is adapted to the available workspace. Training uses bundled Iris data without an external dataset download.")
code("src/        model training and prediction\ndata/       dataset export and provenance\nmodels/     trained model and evaluation\ntests/      automated acceptance checks\ndocs/       Scrum records report and evidence\n.github/    GitHub Actions workflow\ntools/      packaging and evidence helpers")
h("Step 2 Assign Scrum roles")
table(["Role", "Student", "Responsibility"], [
    ["Product Owner", "Divyaprasad Sahu", "Prioritise stories and accept the 95% accuracy goal"],
    ["Scrum Master", "Divyaprasad Sahu", "Facilitate ceremonies and resolve blockers"],
    ["ML Engineer", "Divyaprasad Sahu", "Train reproducibly and save the model"],
    ["Data Engineer", "Divyaprasad Sahu", "Load audit export and split the data"],
    ["Tester", "Divyaprasad Sahu", "Validate accuracy predictions and automation"],
], [1.25, 1.5, 4.1])
h("Tools")
p("Python 3.11 in GitHub Actions; Python 3.12 locally. Pinned scikit-learn 1.7.2, NumPy 2.3.3, SciPy 1.16.2, joblib 1.5.2, threadpoolctl 3.6.0 and pytest 8.4.2. Git manages the history and GitHub hosts the code and workflows.")

page("Product backlog and Sprint planning")
h("Step 3 Create the Product Backlog")
table(["ID", "Story", "Priority", "Acceptance criterion"], [
    ["US1", "Download Iris Dataset", "High", "Load 150 rows and export four features and labels"],
    ["US2", "Clean Dataset", "High", "Validate shape finite positive data and class labels; audit duplicates"],
    ["US3", "Train Decision Tree", "High", "Train 120 rows with seed 42; reproduce predictions"],
    ["US4", "Evaluate Accuracy", "Medium", "At least 95% on 30 held-out rows; save metrics"],
    ["US5", "Save Model", "Medium", "Reload saved model with identical predictions"],
    ["US6", "Prepare Documentation", "Low", "Provide setup Scrum records results and report"],
], [.45, 1.6, .7, 4.1])
p("High-priority work comes first because evaluation depends on valid data and a trained classifier. All six stories remain necessary for the final increment. Additional rubric work covers Git, GitHub, automated tests and delivery.")
h("Step 4 Plan Sprint 1")
p("Sprint duration: one week represented by five simulated working days. Sprint Goal: build and deliver the first working Iris ML model. All six stories are selected.")
table(["Task", "Owner role", "Estimate"], [
    ["Workspace and dependencies", "Scrum Master", "20 min"], ["Load and export data", "Data Engineer", "20 min"],
    ["Audit and prepare features", "Data Engineer", "30 min"], ["Train Decision Tree", "ML Engineer", "30 min"],
    ["Evaluate and automate tests", "Tester", "45 min"], ["Save and reload model", "ML Engineer", "20 min"],
    ["Git and Actions", "Scrum Master", "45 min"], ["Report and demo", "Product Owner", "60 min"],
], [3.55, 2.3, 1])
p("Total estimate: 270 minutes. Estimates are planning values rather than time sheets. Definition of Done: acceptance criteria met, tests passed, code reviewed, main delivery verified and evidence documented.")

page("Daily Scrum simulation")
h("Step 5 Conduct Daily Scrum")
p("These records simulate the intended one-week lifecycle. They do not claim that five people held meetings on five actual calendar days. Each update covers yesterday, today and blockers.")
table(["Day", "Yesterday", "Today", "Blockers and response"], [
    ["1", "Identify goal and roles", "Workspace and backlog", "Check tools and access"],
    ["2", "Prepare requirements", "Audit data and train", "Install missing libraries"],
    ["3", "Prepare baseline", "Test and persist model", "Investigate any failing checks"],
    ["4", "Verify local results", "Connect Actions and propose change", "Inspect access or workflow errors"],
    ["5", "Prepare increment", "Review and retrospective", "Keep acceptance pending until checks pass"],
], [.4, 1.65, 2.15, 2.65])
h("Role updates for simulated day 3")
table(["Role", "Yesterday", "Today", "Blocker"], [
    ["Product Owner", "Set requirements", "Review acceptance evidence", "None simulated"],
    ["Scrum Master", "Plan dependencies", "Facilitate and unblock", "Environment access"],
    ["ML Engineer", "Select Decision Tree", "Train and save", "Validated data needed"],
    ["Data Engineer", "Load Iris", "Validate and split", "Missing dependencies"],
    ["Tester", "Prepare test cases", "Evaluate and run tests", "Trained model needed"],
], [1.2, 1.55, 2.4, 1.7])
h("Actual implementation blockers resolved")
p("The runtime's environment did not bootstrap pip directly, so dependencies were installed into the isolated environment through the bundled runtime's pip. Restricted network access required approved installation and Git pushes. Initial pytest temporary-directory errors were resolved by rerunning with the required environment permissions.")
p("The GitHub connector could not create the PR for the new repository. The authenticated GitHub browser was used instead. Technical evidence records successful final executions separately from the simulation.")

page("Model development and measured results")
h("Step 6 Build the ML model")
p("Load Iris, validate its four positive finite measurements and labels, and split before training. Fit the Decision Tree only on training rows, then compare held-out predictions with their true labels. Fix both random states to reproduce results.")
code("iris = load_iris()\nX_train, X_test, y_train, y_test = train_test_split(\n    iris.data, iris.target, test_size=0.2, random_state=42\n)\nmodel = DecisionTreeClassifier(random_state=42)\nmodel.fit(X_train, y_train)\npredictions = model.predict(X_test)\naccuracy = accuracy_score(y_test, predictions)")
table(["Measurement", "Observed result"], [
    ["Dataset", "150 rows; 4 features; 3 classes with 50 rows each"],
    ["Training and testing", "120 training rows and 30 test rows"],
    ["Missing values", metrics["missing_values"]], ["Duplicate rows", str(metrics["duplicate_rows"]) + " retained for the original split"],
    ["Held-out accuracy", f"{metrics['accuracy']:.2%} = 30 correct predictions / 30"],
    ["Acceptance threshold", "95% achieved"],
], [2, 4.85])
h("Confusion matrix")
table(["Actual class", "Setosa", "Versicolor", "Virginica"],
      [[name, *row] for name, row in zip(metrics["target_names"], metrics["confusion_matrix"])], [2.35, 1.5, 1.5, 1.5])
h("Prediction and persistence")
code("python -m src.iris_demo train\npython -m src.iris_demo predict 5.1 3.5 1.4 0.2\nPrediction: setosa\nLeaf probability: 100.00% (not calibrated confidence)")
p("Training saves a joblib model bundle, metrics JSON, confusion-matrix CSV and data export. Saved-model tests verify identical predictions after reloading. The leaf probability describes training examples in the reached leaf; it does not guarantee correctness. A small single hold-out split is limited evidence of wider performance.")

page("Git and GitHub code management")
h("Repository and change history")
link("Public project", "https://github.com/Divyaprasad-Sahu/mlops-agile-demo")
p("Only the experiment folder is a Git repository. Environments, caches, generated binary models and local temporary reports are excluded. Source, tests, dataset, instructions and workflow are organised and committed with clear messages.")
table(["Command", "Purpose and evidence"], [
    ["git init -b main", "Create the standalone repository"], ["git status", "Inspect pending files before staging"],
    ["git add and git commit", "Save coherent setup model tests automation and documentation changes"],
    ["git log --oneline --graph --all", "Inspect commit history and branch relationships"],
    ["git switch -c feat/prediction-confidence", "Isolate the demonstration feature"],
    ["git diff --stat", "Inspect the proposed feature changes"],
    ["git remote add origin and git push -u", "Share local history with the public GitHub repository"],
    ["git pull --ff-only", "Synchronise main after the reviewed merge"],
], [3.0, 3.85])
h("Actual commit milestones")
code(subprocess.check_output(["git", "log", "--oneline", "-8"], cwd=ROOT, text=True).strip())
h("Demonstration change")
p("The initial predictor returned only the flower name. A feature branch adds the selected class's leaf probability, an acceptance test and documentation. The PR demonstrates how a code change triggers CI before merge and delivery after merge.")
if evidence.get("pr_url"):
    link("Demonstration pull request", evidence["pr_url"])
p("The raw local command transcript is retained in docs/evidence/git_commands.txt. GitHub preserves the corresponding commit and pull-request history.")

page("Automated testing workflow")
h("Acceptance test design")
p("pytest executes the same acceptance suite locally and in GitHub Actions. The tests exercise useful behaviour rather than relying only on printed accuracy.")
table(["Test area", "Acceptance check"], [
    ["Dataset contract", "150 rows four features and balanced class labels"],
    ["Data validation", "Reject NaN negative measurements wrong shape invalid labels and mismatched targets"],
    ["Reproducibility", "Repeated training gives equal metrics and predictions"],
    ["Model quality", "Held-out accuracy at least 95%; 120/30 split and matrix total 30"],
    ["Known predictions", "Representative setosa versicolor and virginica inputs"],
    ["Input validation", "Reject wrong-length zero negative and infinite prediction inputs"],
    ["Save and reload", "Identical predictions after joblib reload; valid reports and CSV export"],
    ["Import behaviour", "Importing the module does not fit a model"],
    ["CLI integration", "Train save predict and explain a missing model"],
    ["Demonstration feature", "Returned leaf probability matches the selected class"],
], [1.75, 5.1])
h("Execution")
code("python -m pip install -r requirements-dev.txt\npython -m pytest -q --junitxml=reports/junit.xml\n19 passed")
p("The final local suite passed all 19 parametrised cases. The workflow retains JUnit XML as a test-results artifact, even on a test failure when a report exists. Exceptions and nonzero command exits fail the job.")
h("Quality gate")
p("A failing test prevents the dependent delivery job. Training also checks the 95% threshold before saving delivery outputs. PR runs validate changes but skip main-only delivery.")

page("Continuous integration and delivery evidence")
h("Pipeline flow")
p("Code change → Git push → pull request checks → successful review and merge → main tests → training and packaging → downloadable model and source increment.")
p("The workflow uses Ubuntu and Python 3.11, installs pinned dependencies, runs pytest, and uploads test results. The delivery job needs a successful test job and is limited to main-branch events outside pull requests.")
h("Actual online evidence")
for label, key in (("Passing PR test run", "pr_run_url"), ("Passing merged main delivery", "main_run_url")):
    if evidence.get(key):
        link(label, evidence[key])
table(["Evidence", "Verified value"], [
    ["Feature commit", evidence.get("feature_sha", "Pending")[:12]],
    ["PR status", evidence.get("pr_status", "Pending")],
    ["Merge commit", evidence.get("merge_sha", "Pending")[:12]],
    ["PR test result", evidence.get("pr_tests", "Pending")],
    ["Main jobs", evidence.get("main_jobs", "Pending")],
    ["Delivery artifact", evidence.get("artifact_name", "Pending")],
    ["Artifact verification", evidence.get("artifact_verified", "Pending")],
], [1.75, 5.1])
h("How to retrieve the increment")
p("Open the passing main run under Actions. Download the iris-delivery artifact and extract the enclosed iris-delivery.zip. It contains the trained model, metrics, source, dataset and usage instructions. Install requirements.txt and run prediction against the included model.")
p("Artifacts are retained for 90 days. The delivery is a downloadable tested package; this experiment does not require a hosted API. Live run links and logs allow the evaluator to connect the source change with testing and delivery.")

page("Sprint review retrospective and increment")
h("Step 7 Review the Sprint")
p("Review the trained model, the 100% held-out result, the example prediction, all 19 passing tests, the PR change and the main delivery artifact. The six backlog stories meet the prepared acceptance criteria, supported by model and workflow evidence.")
h("Step 8 Conduct the retrospective")
p("What went well: a fixed split and seed made results reproducible. Small user stories gave clear dependencies. Invalid-input and persistence tests checked more than model accuracy, while Git history made progress inspectable.")
p("Problems encountered: dependency setup, restricted network and temporary-directory permissions, and limited connector access to the newly created repository. These were resolved using the isolated environment, approved execution and authenticated GitHub browser.")
table(["Improvement", "Owner role", "Check next Sprint"], [
    ["Install libraries and verify tool access before starting", "Scrum Master", "Fresh setup passes"],
    ["Write acceptance tests with each story", "Tester", "Every PR has tests"],
    ["Keep seeds and requirements explicit", "ML Engineer", "Results reproduce"],
    ["Review links and delivered outputs before acceptance", "Product Owner", "Evidence matches accepted commit"],
], [3.2, 1.2, 2.45])
h("Step 9 Deliver the increment")
for text in ("Validated Iris dataset and exported CSV", "Reproducible Decision Tree source and trained model", "Accuracy report confusion matrix and sample prediction", "Automated tests and GitHub Actions workflow", "Meaningful Git history and reviewed demonstration change", "Scrum documentation final report and downloadable delivery"):
    p("Completed  " + text, "List Bullet")
p("Potential next-Sprint work is cross-validation, algorithm comparison and a prediction interface. These are future backlog items, not part of the delivered experiment.")

page("Rubric mapping conclusion and viva")
table(["Category", "Weight", "Evidence"], [
    ["Git Command Usage", "20%", "Actual init status add commit log branch diff and push history"],
    ["GitHub Code Management", "20%", "Public organised repository meaningful commits and reviewed PR"],
    ["Python Project Organisation", "20%", "Separate source data models tests docs and runnable instructions"],
    ["Automated Testing Workflow", "20%", "19 passing tests and uploaded JUnit report"],
    ["CI/CD Demonstration", "20%", "Passing PR checks followed by merged main model delivery"],
], [2.05, .6, 4.2])
h("Conclusion")
p("Agile and Scrum were demonstrated through an individual Iris MLOps simulation covering roles, Product Backlog, Sprint planning, Daily Scrum, review, retrospective and a working increment. The project meets the 95% model acceptance criterion on the specified split and connects Git-managed changes to automated tests and delivery. The evidence matrix targets the Excellent descriptors; the evaluator assigns the final grade.")
h("Viva preparation")
p("Product Backlog versus Sprint Backlog: product-wide ordered work versus the Sprint Goal, selected work and execution plan. Review versus retrospective: inspect the product increment versus improve how the team works. CI versus CD: validate changes automatically versus deliver the accepted software package.")
p("Why hold out data: estimate performance on examples excluded from model fitting. Why seed 42: repeat the split and training behaviour. Why save the model: predict without retraining and verify delivery. Why tests matter: detect regressions before the increment is accepted.")
h("References and supporting records")
p("Course handout: Experiment 1 Agile and Scrum Demo. Detailed Scrum records, Product_Backlog.md, Demonstration_Guide.md, Rubric_Mapping.md and actual evidence are included in the repository.")
link("Scrum Guide", "https://scrumguides.org/scrum-guide.html")
link("Iris dataset documentation", "https://scikit-learn.org/stable/modules/generated/sklearn.datasets.load_iris.html")
link("GitHub Actions", "https://docs.github.com/en/actions")

page("Code excerpts for VS Code")
p("The following excerpts are copied directly from the submitted project files. Filenames and actual line numbers make each section easy to locate in VS Code. The line numbers are display labels and are not part of the Python code.")
h("Dataset validation")
function_excerpt("src/iris_demo.py", "validate_dataset")
h("Reproducible training and evaluation")
function_excerpt("src/iris_demo.py", "train_model")
p("The split and classifier use seed 42. Evaluation uses only held-out examples, and the resulting metrics include accuracy, class labels, a confusion matrix and a dataset audit.")

page("Prediction and saved model code")
h("Validated prediction with leaf probability")
function_excerpt("src/iris_demo.py", "prediction_details")
p("Four finite positive measurements are required. The selected class and its leaf probability are returned together; the CLI explicitly states that this probability is not calibrated confidence.")
h("Model persistence and dataset export")
function_excerpt("src/iris_demo.py", "save_training_outputs")
p("joblib saves the model bundle. JSON and CSV preserve evaluation results, and the Iris export makes the dataset inspectable. The automated round-trip test verifies the saved model's predictions.")

page("Automated test code excerpts")
h("Reproducibility and accuracy acceptance test")
function_excerpt("tests/test_iris_demo.py", "test_training_is_reproducible_and_meets_acceptance")
p("This test verifies the 95% threshold, the 120/30 split, the confusion-matrix total and agreement between repeated model predictions.")
h("Saved model round trip test")
function_excerpt("tests/test_iris_demo.py", "test_saved_model_and_reports_roundtrip")
p("The model is saved in an isolated temporary directory, loaded again, and checked against its original predictions. The report and exported dataset are also verified.")
code("python -m pytest -q --junitxml=reports/junit.xml\n19 passed")

page("GitHub Actions workflow excerpts")
h("Automated test job")
source_excerpt(".github/workflows/ci.yml", 11, 32)
p("Every push and pull request runs this job with Python 3.11. The JUnit test-results artifact is uploaded even when tests fail, provided the report exists.")
h("Dependent delivery job")
source_excerpt(".github/workflows/ci.yml", 34, 57)
p("The needs field gates delivery on successful tests. The condition restricts delivery to main, where training and packaging produce the downloadable accepted increment.")

footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run("Experiment 1   |   ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(9)
target = ROOT / "docs/Experiment_1_Agile_and_Scrum_Report.docx"
doc.save(target)
print(target)
